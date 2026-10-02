"""Atlanta Fed BIE connector. Fixture: the live 2026-10-01 workbook trimmed to
the title/banner/header rows + the last 110 monthly rows of "BIE Survey
results" and the full quarterly long-term sheet (values untouched)."""
import io
from pathlib import Path

import openpyxl
import pytest

from pipeline.connectors import atlfed

XLSX = (Path(__file__).parent / "fixtures" / "atlfed_bie.xlsx").read_bytes()


class _B:
    def __init__(self, content):
        self.content = content

    def raise_for_status(self):
        pass


def _get(content=XLSX):
    return lambda url, timeout=None: _B(content)


def test_reads_year_ahead_median_and_long_term_mean_by_survey_month():
    obs = atlfed.fetch(["BIE_1Y_MEDIAN", "BIE_LT_MEAN"], vintage_date="2026-10-01", http_get=_get())
    yr = [o for o in obs if o.series_code == "atl_bie_1y_median"]
    lt = [o for o in obs if o.series_code == "atl_bie_lt_mean"]
    # released 2026-09-19 -> stored as the survey month, 0.0236 -> 2.36%
    assert (yr[-1].obs_date, yr[-1].value) == ("2026-09-01", 2.36)
    assert (lt[-1].obs_date, lt[-1].value) == ("2026-09-01", 2.775)
    assert len({o.obs_date for o in yr}) == len(yr)  # one row per month
    assert {(o.source, o.route) for o in obs} == {("ATLFED", "XLSX")}


def _mutate(fn):
    wb = openpyxl.load_workbook(io.BytesIO(XLSX))
    fn(wb)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_moved_question_banner_is_drift():
    def rename(wb):
        for row in wb["BIE Survey results"].iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.startswith("Question 4"):
                    c.value = "Q4 (renamed)"
    with pytest.raises(ValueError, match="structure drift"):
        atlfed.fetch(["BIE_1Y_MEDIAN"], http_get=_get(_mutate(rename)))


def test_long_term_sheet_missing_is_partial_not_fatal():
    import warnings
    from pipeline.connectors.util import PartialFetchWarning
    content = _mutate(lambda wb: wb.remove(wb["Quarterly - Long-Term Infl Exp"]))
    with warnings.catch_warnings(record=True) as caught:
        warnings.simplefilter("always")
        obs = atlfed.fetch(["BIE_1Y_MEDIAN", "BIE_LT_MEAN"], http_get=_get(content))
    assert {o.series_code for o in obs} == {"atl_bie_1y_median"}
    assert any(issubclass(w.category, PartialFetchWarning) for w in caught)


def test_implausible_value_is_drift():
    def blow_up(wb):
        ws = wb["BIE Survey results"]
        ws.cell(row=ws.max_row, column=31).value = 0.9  # 90%
    with pytest.raises(ValueError, match="implausible"):
        atlfed.fetch(["BIE_1Y_MEDIAN"], http_get=_get(_mutate(blow_up)))
