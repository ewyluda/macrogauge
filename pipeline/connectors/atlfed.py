"""Atlanta Fed Business Inflation Expectations (BIE) survey — firms'
year-ahead unit-cost expectations (monthly) and 5-10-year expectations
(quarterly).

Keyless public xlsx (verified live 2026-10-01; the old
/research/inflationproject/bie URL redirects here). Sheet "BIE Survey
results": row 1 carries question banners ("Question 4: Projecting ahead ...
unit costs over the next 12 months."), row 2 the per-question column headers
(Mean / Median / ...), data rows start with the survey RELEASE date (e.g.
2026-09-19). Sheet "Quarterly - Long-Term Infl Exp": header row "Month",
"Number of Responses", ..., "Mean". Values are fractions (0.0236 = 2.36%).

Stored by survey month (the release date's month: the survey is fielded and
published within the same month), with the release date in vintage_date —
the repo-wide convention that obs_date is the period the data describes.
Unofficial-endpoint treatment: question banner, header strings, date form
and a plausible range are pinned; any change raises "structure drift?".
"""
import io
import re
from datetime import date, datetime

import requests

from pipeline.connectors.fred import today_et
from pipeline.connectors.util import get_bytes, warn_partial
from pipeline.models import Observation

URL = ("https://www.atlantafed.org/-/media/Project/Atlanta/FRBA/Documents/"
       "research-and-data/data/bie/data/bie.xlsx")
_UA = ("Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
       "(KHTML, like Gecko) Chrome/124.0 Safari/537.36")
PLAUSIBLE = (-5.0, 15.0)   # % expected unit-cost growth
Q4_BANNER = "Question 4: Projecting ahead"
SERIES = {"BIE_1Y_MEDIAN": "atl_bie_1y_median", "BIE_LT_MEAN": "atl_bie_lt_mean"}


def _default_get(url, timeout=60):
    return requests.get(url, timeout=timeout, headers={"User-Agent": _UA})


def _month(cell) -> str:
    if isinstance(cell, (datetime, date)):
        d = cell
    elif isinstance(cell, str) and cell.strip():
        try:
            d = datetime.strptime(cell.strip(), "%B %d, %Y")
        except ValueError:
            raise ValueError(f"atlfed bie: date {cell!r} not a date (structure drift?)") from None
    else:
        raise ValueError(f"atlfed bie: date {cell!r} not a date (structure drift?)")
    return f"{d.year:04d}-{d.month:02d}-01"


def _pct(v, where: str) -> float:
    pct = float(v) * 100
    if not PLAUSIBLE[0] <= pct <= PLAUSIBLE[1]:
        raise ValueError(f"atlfed bie {where}: {pct:.2f}% implausible (structure drift?)")
    return round(pct, 3)


def _year_ahead_median(wb, vintage: str) -> list[Observation]:
    rows = list(wb["BIE Survey results"].iter_rows(values_only=True))
    banner_row = next((i for i, r in enumerate(rows[:6])
                       if any(isinstance(c, str) and c.startswith(Q4_BANNER) for c in r)), None)
    if banner_row is None:
        raise ValueError("atlfed bie: Question 4 banner not found (structure drift?)")
    start = next(i for i, c in enumerate(rows[banner_row])
                 if isinstance(c, str) and c.startswith(Q4_BANNER))
    nxt = next((i for i in range(start + 1, len(rows[banner_row]))
                if isinstance(rows[banner_row][i], str)
                and re.match(r"Question \d+:", rows[banner_row][i])), len(rows[banner_row]))
    header = rows[banner_row + 1]
    col = next((i for i in range(start, nxt) if header[i] == "Median"), None)
    if col is None:
        raise ValueError("atlfed bie: Question 4 'Median' column missing (structure drift?)")
    out = []
    for r in rows[banner_row + 2:]:
        if not r or r[0] is None or col >= len(r) or r[col] is None:
            continue
        m = _month(r[0])
        out.append(Observation("atl_bie_1y_median", m, _pct(r[col], m), vintage, "ATLFED", "XLSX"))
    if len(out) < 100:
        raise ValueError(f"atlfed bie: only {len(out)} monthly rows (structure drift?)")
    return out


def _long_term_mean(wb, vintage: str) -> list[Observation]:
    name = "Quarterly - Long-Term Infl Exp"
    if name not in wb.sheetnames:
        raise ValueError(f"atlfed bie: sheet {name!r} missing (structure drift?)")
    rows = list(wb[name].iter_rows(values_only=True))
    head = next((i for i, r in enumerate(rows[:8]) if r and r[0] == "Month"), None)
    if head is None or "Mean" not in rows[head]:
        raise ValueError("atlfed bie: long-term header changed (structure drift?)")
    col = rows[head].index("Mean")
    out = []
    for r in rows[head + 1:]:
        if not r or r[0] is None or r[col] is None:
            continue
        m = _month(r[0])
        out.append(Observation("atl_bie_lt_mean", m, _pct(r[col], m), vintage, "ATLFED", "XLSX"))
    if len(out) < 40:
        raise ValueError(f"atlfed bie: only {len(out)} quarterly rows (structure drift?)")
    return out


def fetch(source_ids: list[str], vintage_date: str | None = None,
          http_get=None) -> list[Observation]:
    """source_ids: BIE_1Y_MEDIAN and/or BIE_LT_MEAN — one workbook; each
    series fails independently (partial warning); all failing raises."""
    import openpyxl
    http_get = http_get or _default_get
    vintage = vintage_date or today_et()
    wb = openpyxl.load_workbook(io.BytesIO(get_bytes(URL, http_get)),
                                read_only=True, data_only=True)
    if "BIE Survey results" not in wb.sheetnames:
        raise ValueError("atlfed bie: 'BIE Survey results' sheet missing (structure drift?)")
    readers = {"BIE_1Y_MEDIAN": _year_ahead_median, "BIE_LT_MEAN": _long_term_mean}
    out, errors = [], []
    for sid in source_ids or list(SERIES):
        try:
            out.extend(readers[sid](wb, vintage))
        except Exception as e:  # per-series isolation
            errors.append((sid, e))
    if errors and not out:
        raise errors[0][1]
    warn_partial("ATLFED", errors)
    return out
