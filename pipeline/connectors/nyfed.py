"""NY Fed Global Supply Chain Pressure Index (GSCPI) — monthly, std devs.

Keyless public CSV behind the NY Fed's GSCPI interactive (verified live
2026-09-28). The advertised download, .../gscpi/downloads/gscpi_data.xlsx,
is a legacy BIFF .xls despite its extension (openpyxl cannot read it, and
its file metadata was last saved 2024-11); the interactive's own data file
is plain CSV and is a VINTAGE MATRIX: one row per reference month
("30-Sep-1997" = month-end), one column per publication vintage ("Jan-22"
.. "Sep-26"), "#N/A" where a vintage had not yet covered the month. The
LAST column is the current publication — the only one read here; earlier
columns are the index's own revision history.

Treated as an unofficial endpoint (it is an interactive's data file, not a
documented download): the header shape, the vintage-label pattern, the
date format and a plausible-value range are all pinned, and any deviation
raises "structure drift?" into collect isolation.
"""
import csv
import io
import re
from datetime import datetime

import requests

from pipeline.connectors.fred import today_et
from pipeline.connectors.util import warn_partial
from pipeline.models import Observation

URL = ("https://www.newyorkfed.org/medialibrary/research/interactives/data/"
       "gscpi/gscpi_interactive_data.csv")
_UA = ("Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
       "(KHTML, like Gecko) Chrome/124.0 Safari/537.36")
VINTAGE_RE = re.compile(r"^[A-Z][a-z]{2}-\d{2}$")      # "Sep-26"
PLAUSIBLE = (-5.0, 10.0)    # std devs from average (record: ~4.3, Dec 2021)
MIN_MONTHS = 300            # the series starts 1997-09 (~350 months by 2026)


def _default_get(url, timeout=60):
    return requests.get(url, timeout=timeout, headers={"User-Agent": _UA})


def fetch(source_ids: list[str], vintage_date: str | None = None,
          http_get=None) -> list[Observation]:
    """source_ids: "GSCPI" and/or "MCT". Each series fails independently
    (a partial warning), so a redesigned MCT chart file never takes GSCPI
    down; every series failing raises."""
    http_get = http_get or _default_get
    vintage = vintage_date or today_et()
    out, errors = [], []
    ids = source_ids or ["GSCPI"]
    for sid in [i for i in ids if not i.startswith("SCE_")]:
        try:
            out.extend((_fetch_mct if sid == "MCT" else _fetch_gscpi)(sid, vintage, http_get))
        except Exception as e:  # per-series isolation
            errors.append((sid, e))
    sce_ids = [i for i in ids if i.startswith("SCE_")]
    if sce_ids:  # one workbook serves every SCE series
        try:
            out.extend(_fetch_sce(sce_ids, vintage, http_get))
        except Exception as e:
            errors.append(("SCE", e))
    if errors and not out:
        raise errors[0][1]
    warn_partial("NYFED", errors)
    return out


SCE_URL = ("https://www.newyorkfed.org/medialibrary/interactives/sce/sce/downloads/data/"
           "frbny-sce-data.xlsx")
# source_id -> (sheet, exact column header, store series code)
SCE_SERIES = {
    "SCE_1Y": ("Inflation expectations", "Median one-year ahead expected inflation rate", "nyfed_sce_1y"),
    "SCE_3Y": ("Inflation expectations", "Median three-year ahead expected inflation rate", "nyfed_sce_3y"),
    "SCE_5Y": ("Five-year ahead Infl Exp", "Median five-year ahead expected inflation rate", "nyfed_sce_5y"),
}
SCE_PLAUSIBLE = (-5.0, 20.0)  # % expected inflation (1y median peaked ~6.8 in 2022)


def _fetch_sce(source_ids: list[str], vintage: str, http_get) -> list[Observation]:
    """NY Fed Survey of Consumer Expectations medians (verified live
    2026-10-01): one xlsx, a title block, then a header row whose first cell
    is empty and whose cells are the series names; data rows start with the
    survey month as an integer YYYYMM (202608). Stored by survey month (the
    month the data describes), the release date lives in vintage_date — the
    repo convention. The full history republishes monthly; the store's
    value-dedupe keeps only the changes. Sheet names, the exact header strings,
    the YYYYMM form and a plausible range are pinned."""
    import openpyxl
    from pipeline.connectors.util import get_bytes
    wb = openpyxl.load_workbook(io.BytesIO(get_bytes(SCE_URL, http_get)),
                                read_only=True, data_only=True)
    out = []
    for sid in source_ids:
        if sid not in SCE_SERIES:
            raise ValueError(f"nyfed sce: unknown source_id {sid!r}")
        sheet, header, code = SCE_SERIES[sid]
        if sheet not in wb.sheetnames:
            raise ValueError(f"nyfed sce: sheet {sheet!r} missing (structure drift?)")
        rows = list(wb[sheet].iter_rows(values_only=True))
        head = next((i for i, r in enumerate(rows[:12]) if r and header in r), None)
        if head is None:
            raise ValueError(f"nyfed sce: header {header!r} not found (structure drift?)")
        col = rows[head].index(header)
        series = []
        for r in rows[head + 1:]:
            if not r or r[0] is None:
                continue
            ym = str(r[0]).strip()
            if not re.fullmatch(r"(19|20)\d{2}(0[1-9]|1[0-2])", ym):
                raise ValueError(f"nyfed sce: month {r[0]!r} is not YYYYMM (structure drift?)")
            v = r[col] if col < len(r) else None
            if v is None:
                continue
            v = float(v)
            if not SCE_PLAUSIBLE[0] <= v <= SCE_PLAUSIBLE[1]:
                raise ValueError(f"nyfed sce {sid} {ym}: {v} implausible (structure drift?)")
            series.append(Observation(series_code=code, obs_date=f"{ym[:4]}-{ym[4:]}-01",
                                      value=round(v, 4), vintage_date=vintage,
                                      source="NYFED", route="XLSX"))
        if len(series) < 24:
            raise ValueError(f"nyfed sce {sid}: only {len(series)} months (structure drift?)")
        out.extend(series)
    return out


MCT_URL = ("https://www.newyorkfed.org/medialibrary/Research/Interactives/Data/"
           "mct/mct-chart-data.csv")
MCT_PLAUSIBLE = (-2.0, 12.0)   # % annual trend inflation (1970s peak ~8)
MCT_MIN_MONTHS = 600           # starts 1960-01


def _fetch_mct(series_code: str, vintage: str, http_get) -> list[Observation]:
    """NY Fed Multivariate Core Trend (PCE-based trend inflation, %). The
    interactive's chart CSV (verified live 2026-09-29) carries three header
    rows, then a 'column name' row: Date, then four "MCT" columns — 68% band
    low, MCT point estimate, band high, normalized MCT — then headline and
    core PCE YoY and the sector decomposition. Only the point estimate (the
    second MCT column) is read; the header shape is pinned."""
    resp = http_get(MCT_URL, timeout=60)
    resp.raise_for_status()
    rows = list(csv.reader(io.StringIO(resp.text.lstrip("\ufeff"))))
    head = next((i for i, r in enumerate(rows) if r and r[0].strip() == "column name"), None)
    if head is None:
        raise ValueError("nyfed mct: no 'column name' header row (structure drift?)")
    h = [c.strip() for c in rows[head]]
    if h[1:6] != ["Date", "MCT", "MCT", "MCT", "MCT"] or "Core PCE inflation (YoY)" not in h:
        raise ValueError(f"nyfed mct: header {h[:8]!r} changed (structure drift?)")
    out = []
    for r in rows[head + 1:]:
        if len(r) < 4 or not r[1].strip():
            continue
        try:
            month = datetime.strptime(r[1].strip(), "%m/%d/%Y")
        except ValueError:
            raise ValueError(f"nyfed mct: date {r[1]!r} is not M/D/YYYY (structure drift?)") from None
        low, mct, high = (float(r[2]), float(r[3]), float(r[4]))
        if not (MCT_PLAUSIBLE[0] <= mct <= MCT_PLAUSIBLE[1]) or not (low <= mct <= high):
            raise ValueError(f"nyfed mct {r[1]}: {mct} (band {low}-{high}) implausible (structure drift?)")
        out.append(Observation(series_code="nyfed_mct", obs_date=f"{month.year}-{month.month:02d}-01",
                               value=mct, vintage_date=vintage, source="NYFED", route="CSV"))
    if len(out) < MCT_MIN_MONTHS:
        raise ValueError(f"nyfed mct: only {len(out)} months (< {MCT_MIN_MONTHS}; structure drift?)")
    return out


def _fetch_gscpi(series_code: str, vintage: str, http_get) -> list[Observation]:
    resp = http_get(URL, timeout=60)
    resp.raise_for_status()
    rows = list(csv.reader(io.StringIO(resp.text)))
    if not rows or not rows[0] or rows[0][0].strip() != "Date":
        raise ValueError("nyfed gscpi: header does not start with 'Date' "
                         "(structure drift?)")
    header = [h.strip() for h in rows[0]]
    col = len(header) - 1
    while col > 0 and not header[col]:
        col -= 1
    label = header[col]
    if col < 1 or not VINTAGE_RE.match(label):
        raise ValueError(f"nyfed gscpi: last vintage column {label!r} is not "
                         "'Mon-YY' (structure drift?)")
    pub = datetime.strptime(label, "%b-%y")
    out = []
    for r in rows[1:]:
        if not r or not r[0].strip():
            continue                      # trailing blank row(s)
        try:
            month = datetime.strptime(r[0].strip(), "%d-%b-%Y")
        except ValueError:
            raise ValueError(f"nyfed gscpi: date {r[0]!r} is not DD-Mon-YYYY "
                             "(structure drift?)") from None
        cell = r[col].strip() if col < len(r) else ""
        if cell in ("", "#N/A"):
            continue
        value = float(cell)
        if not PLAUSIBLE[0] <= value <= PLAUSIBLE[1]:
            raise ValueError(f"nyfed gscpi {r[0]}: {value} outside {PLAUSIBLE} "
                             "(structure drift?)")
        if (month.year, month.month) >= (pub.year, pub.month):
            raise ValueError(f"nyfed gscpi: vintage {label} covers {r[0]} — a "
                             "publication cannot cover its own month (structure drift?)")
        out.append(Observation(series_code=series_code,
                               obs_date=f"{month.year}-{month.month:02d}-01",
                               value=value, vintage_date=vintage,
                               source="NYFED", route="CSV"))
    if len(out) < MIN_MONTHS:
        raise ValueError(f"nyfed gscpi: only {len(out)} months in vintage {label} "
                         f"(< {MIN_MONTHS}; structure drift?)")
    return out
